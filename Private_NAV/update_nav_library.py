"""Validate screenshot-extracted NAVs, update the NAV workbook, and export this batch."""

from __future__ import annotations

import argparse
import csv
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parent
LIBRARY = ROOT / "私募净值.xlsx"
MASTER = ROOT / "私募名单.xlsx"
REQUIRED = {"source_file", "observed_name", "matched_name", "date", "unit_nav", "cumulative_nav", "status", "notes"}


def read_batch(path: Path) -> List[dict]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        missing = REQUIRED - set(reader.fieldnames or [])
        if missing:
            raise ValueError("批次文件缺少列: {}".format(", ".join(sorted(missing))))
        rows = []
        for line, row in enumerate(reader, 2):
            row["line"] = line
            row["status"] = row["status"].strip().lower()
            if row["status"] == "accepted":
                row["date"] = datetime.strptime(row["date"], "%Y-%m-%d").date()
                row["unit_nav"] = float(row["unit_nav"])
                row["cumulative_nav"] = float(row["cumulative_nav"])
            rows.append(row)
    return rows


def product_codes() -> Dict[str, str]:
    sheet = load_workbook(MASTER, data_only=True).active
    header = {cell: index for index, cell in enumerate(next(sheet.iter_rows(values_only=True)))}
    return {
        str(row[header["基金名称"]]).strip(): "" if row[header["备案编码"]] is None else str(row[header["备案编码"]]).strip()
        for row in sheet.iter_rows(values_only=True)
        if row[header["基金名称"]]
    }


def existing_dates(book) -> Dict[str, Dict[object, Tuple[float, float]]]:
    result = {}
    for sheet in book.worksheets:
        header = {cell: index for index, cell in enumerate(next(sheet.iter_rows(values_only=True)))}
        values = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[header["日期"]] is not None:
                date = row[header["日期"]]
                date = date.date() if hasattr(date, "date") else date
                values[date] = (float(row[header["单位净值"]]), float(row[header["累计净值"]]))
        result[sheet.title] = values
    return result


def classify(rows: List[dict], codes: Dict[str, str], existing: Dict[str, Dict[object, Tuple[float, float]]]):
    additions, review, seen = [], [], {}
    for row in rows:
        product = row["matched_name"].strip()
        if row["status"] != "accepted":
            row["result"] = "未入库: {}".format(row["status"] or "待复核")
            review.append(row)
            continue
        key = (product, row["date"])
        value = (row["unit_nav"], row["cumulative_nav"])
        if product not in codes:
            row["result"] = "未入库: 产品不在名单"
        elif key in seen and seen[key] != value:
            row["result"] = "未入库: 批次内同日期数值冲突"
        elif product in existing and row["date"] in existing[product]:
            old = existing[product][row["date"]]
            row["result"] = "跳过: 已存在相同记录" if old == value else "未入库: 同日期数值冲突"
        else:
            row["备案编码"] = codes[product]
            row["result"] = "新增产品工作表" if product not in existing else "新增"
            additions.append(row)
            seen[key] = value
            continue
        review.append(row)
    return sorted(additions, key=lambda row: (row["matched_name"], row["date"])), review


def update_library(additions: List[dict]) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = ROOT / "backups"
    backup_dir.mkdir(exist_ok=True)
    backup = backup_dir / "私募净值_{}.xlsx".format(stamp)
    shutil.copy2(LIBRARY, backup)
    book = load_workbook(LIBRARY)
    grouped = defaultdict(list)
    for row in additions:
        grouped[row["matched_name"]].append(row)
    for product, records in grouped.items():
        if product not in book.sheetnames:
            sheet = book.create_sheet(product)
            sheet.append(["日期", "单位净值", "累计净值"])
        else:
            sheet = book[product]
        values = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0] is not None:
                value_date = row[0].date() if hasattr(row[0], "date") else row[0]
                values[value_date] = (row[1], row[2])
        for row in records:
            values[row["date"]] = (row["unit_nav"], row["cumulative_nav"])
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
        for value_date, value in sorted(values.items()):
            sheet.append([value_date, value[0], value[1]])
    book.save(LIBRARY)
    return backup


def write_report(additions: List[dict], review: List[dict]) -> Path:
    output_dir = ROOT / "output"
    output_dir.mkdir(exist_ok=True)
    path = output_dir / "新增净值_{}.xlsx".format(datetime.now().strftime("%Y%m%d_%H%M%S"))
    book = Workbook()
    data = book.active
    data.title = "新增净值"
    data.append(["净值日期", "基金名称", "备案编码", "单位净值", "累计净值"])
    for row in additions:
        data.append([row["date"], row["matched_name"], row["备案编码"], row["unit_nav"], row["cumulative_nav"]])
    review_sheet = book.create_sheet("待复核")
    review_sheet.append(["截图文件", "截图产品名", "候选产品", "状态", "说明"])
    for row in review:
        review_sheet.append([row.get("source_file"), row.get("observed_name"), row.get("matched_name"), row.get("result"), row.get("notes")])
    for sheet in book.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[chr(64 + column)].width = 22
    for row in data.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        row[3].number_format = "0.0000"
        row[4].number_format = "0.0000"
    book.save(path)
    return path


def main() -> int:
    parser = argparse.ArgumentParser(description="将截图提取后的净值更新入库")
    parser.add_argument("--batch", required=True, type=Path)
    parser.add_argument("--apply", action="store_true", help="实际更新私募净值.xlsx；否则只生成报告")
    args = parser.parse_args()
    try:
        rows = read_batch(args.batch)
        codes = product_codes()
        existing = existing_dates(load_workbook(LIBRARY, data_only=True))
        additions, review = classify(rows, codes, existing)
        report = write_report(additions, review)
        if args.apply and additions:
            backup = update_library(additions)
            print("净值库已更新，备份: {}".format(backup))
        print("新增 {} 条，待复核 {} 条。".format(len(additions), len(review)))
        print("本次汇总: {}".format(report))
        return 0
    except Exception as exc:
        print("处理失败: {}".format(exc))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
