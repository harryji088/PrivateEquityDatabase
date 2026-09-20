"""Create a sendable NAV-window workbook from successful update reports."""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parent
LIBRARY = ROOT / "私募净值.xlsx"
HEADERS = ["净值日期", "基金名称", "备案编码", "单位净值", "累计净值"]


def current_values() -> Dict[Tuple[str, object], Tuple[float, float]]:
    values = {}
    book = load_workbook(LIBRARY, data_only=True)
    for sheet in book.worksheets:
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            date = row[0].date() if hasattr(row[0], "date") else row[0]
            values[(sheet.title, date)] = (float(row[1]), float(row[2]))
    return values


def collect(reports: List[Path]) -> List[tuple]:
    library = current_values()
    rows: Dict[Tuple[str, object], tuple] = {}
    for report in reports:
        sheet = load_workbook(report, data_only=True)["新增净值"]
        if [cell.value for cell in sheet[1]] != HEADERS:
            raise ValueError("汇总文件列不符合预期: {}".format(report))
        for date, name, code, unit_nav, cumulative_nav in sheet.iter_rows(min_row=2, values_only=True):
            if date is None:
                continue
            date = date.date() if hasattr(date, "date") else date
            key = (str(name), date)
            value = (float(unit_nav), float(cumulative_nav))
            if library.get(key) != value:
                raise ValueError("当前净值库与汇总不一致: {} {}".format(name, date))
            record = (date, str(name), "" if code is None else str(code), value[0], value[1])
            if key in rows and rows[key] != record:
                raise ValueError("窗口内同产品同日期数值冲突: {} {}".format(name, date))
            rows[key] = record
    return sorted(rows.values(), key=lambda row: (row[0], row[1]))


def write_output(rows: List[tuple], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "新增净值"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column, width in zip("ABCDE", [14, 34, 14, 14, 14]):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        row[3].number_format = "0.0000"
        row[4].number_format = "0.0000"
    book.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="生成仅含实际新增净值的邮件窗口附件")
    parser.add_argument("--report", action="append", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    rows = collect(args.report)
    write_output(rows, args.output)
    print("窗口附件已生成：{} 条新增净值，{}".format(len(rows), args.output))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
